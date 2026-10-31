"""Exporta um resultado de análise para uma planilha Excel (.xlsx) com duas
abinas: 'Resumo' (campos escalares) e 'Dados' (tabela), usando a mesma divisão
de `tabela_e_resumo` que os outros renderizadores.
Aplica estilos e formatação premium usando openpyxl."""
import io
import re
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .shape import tabela_e_resumo


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", str(name or "Dados")).strip() or "Dados"
    cleaned = cleaned[:31]
    candidate = cleaned
    idx = 2
    while candidate in used:
        suffix = f"_{idx}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        idx += 1
    used.add(candidate)
    return candidate


def _summary_rows(tool: str, result: Any) -> list[dict[str, Any]]:
    _, resumo = tabela_e_resumo(tool, result)
    rows = []
    for key, value in resumo.items():
        if isinstance(value, (dict, list)) or value is None or key == "markdown":
            continue
        rows.append({"fonte": tool, "metrica": key, "valor": value})
    return rows


def export_xlsx(tool: str, result: Any) -> bytes:
    tabela, resumo = tabela_e_resumo(tool, result)
    resumo_escalar = {k: v for k, v in resumo.items() if not isinstance(v, (dict, list))}

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # 1. Escreve a aba Resumo
        df_resumo = pd.DataFrame(list(resumo_escalar.items()), columns=["Métrica", "Valor"])
        df_resumo.to_excel(writer, sheet_name="Resumo", index=False)
        
        # 2. Escreve a aba Dados
        if tabela:
            df_dados = pd.DataFrame(tabela)
            df_dados.to_excel(writer, sheet_name="Dados", index=False)
        else:
            df_dados = None

        # 3. Aplica estilização via openpyxl
        workbook = writer.book
        
        # Definição de estilos
        font_family = "Segoe UI"
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        cell_font = Font(name=font_family, size=11)
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        thin_border_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            ws.sheet_view.showGridLines = True
            
            # Formata cabeçalho (linha 1)
            ws.row_dimensions[1].height = 28
            for col_num in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center
                cell.border = cell_border

            # Formata o corpo da tabela
            for row_num in range(2, ws.max_row + 1):
                ws.row_dimensions[row_num].height = 20
                is_zebra = (row_num % 2 == 0)
                
                for col_num in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.font = cell_font
                    cell.border = cell_border
                    if is_zebra:
                        cell.fill = zebra_fill
                    
                    # Alinhamento e formatação baseados no tipo de dado
                    val = cell.value
                    col_name = str(ws.cell(row=1, column=col_num).value).lower()
                    
                    if isinstance(val, (int, float)):
                        cell.alignment = align_right
                        
                        # Formatações numéricas específicas
                        if "pct" in col_name or "taxa" in col_name or "conclusao" in col_name or "preenchimento" in col_name or "sla" in col_name:
                            # Se for porcentagem maior que 1 (ex: 51.1 representando 51.1%), usa formato literal %
                            # Se for menor que 1 (ex: 0.511), usa 0.0% padrão
                            if abs(val) > 1.0:
                                cell.number_format = '0.0"%"'
                            else:
                                cell.number_format = '0.0%'
                        elif "valor" in col_name or "score" in col_name or "nps" in col_name or "csat" in col_name or "ces" in col_name or "horas" in col_name:
                            cell.number_format = '0.00'
                        else:
                            cell.number_format = '#,##0'
                    else:
                        cell.alignment = align_left

            # Auto-ajuste de largura de colunas
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    # Se tiver quebra de linha, pega a maior linha
                    lines = val_str.split('\n')
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)
                # Margem de segurança de 4 caracteres, mínimo 12
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return buffer.getvalue()


def export_xlsx_multi(sources: list[dict[str, Any]], title: str = "Analise integrada") -> bytes:
    buffer = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary_rows: list[dict[str, Any]] = []
        for idx, item in enumerate(sources, start=1):
            tool = str(item.get("tool") or f"fonte_{idx}")
            result = item.get("result")
            tabela, resumo = tabela_e_resumo(tool, result)
            summary_rows.extend(_summary_rows(tool, result))

            if tabela:
                df = pd.DataFrame(tabela)
            else:
                resumo_escalar = {
                    k: v for k, v in resumo.items()
                    if not isinstance(v, (dict, list)) and v is not None and k != "markdown"
                }
                df = pd.DataFrame(list(resumo_escalar.items()), columns=["campo", "valor"])
            if df.empty:
                df = pd.DataFrame([{"status": "sem dados exportaveis"}])
            df.to_excel(writer, sheet_name=_safe_sheet_name(tool, used), index=False)

        if summary_rows:
            pd.DataFrame(summary_rows).to_excel(writer, sheet_name=_safe_sheet_name("Resumo Integrado", used), index=False)
        else:
            pd.DataFrame([{"titulo": title, "status": "sem resumo"}]).to_excel(
                writer,
                sheet_name=_safe_sheet_name("Resumo Integrado", used),
                index=False,
            )

        workbook = writer.book
        font_family = "Segoe UI"
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell_font = Font(name=font_family, size=11)
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            ws.sheet_view.showGridLines = True
            ws.freeze_panes = "A2"
            ws.row_dimensions[1].height = 28
            for col_num in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center
                cell.border = cell_border
            for row_num in range(2, ws.max_row + 1):
                for col_num in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.font = cell_font
                    cell.border = cell_border
                    if row_num % 2 == 0:
                        cell.fill = zebra_fill
                    cell.alignment = align_right if isinstance(cell.value, (int, float)) else align_left
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 42))

    return buffer.getvalue()

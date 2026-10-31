"""Tabela para planilha xlsx formatada, sem escrever código de planilha.

    python tabela_para_xlsx.py vendas.csv vendas.xlsx
    python tabela_para_xlsx.py dados.json relatorio.xlsx --aba "Setembro"

Entrada: CSV com cabeçalho; JSON com uma lista de objetos; ou JSON com um
objeto em que cada chave é uma aba e o valor é a lista dela. Sai com cabeçalho
destacado, filtro, primeira linha congelada, largura ajustada e número
gravado como número (soma e fórmula funcionam). Precisa do openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_NUMERO = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+([.,]\d+)?$")


def ler_abas(caminho: Path, aba: str) -> Dict[str, List[Dict[str, Any]]]:
    texto = caminho.read_text(encoding="utf-8-sig")
    if caminho.suffix.lower() == ".json":
        dados = json.loads(texto)
        if isinstance(dados, dict) and all(isinstance(v, list) for v in dados.values()):
            return {str(k): [d for d in v if isinstance(d, dict)] for k, v in dados.items()}
        return {aba: [d for d in (dados if isinstance(dados, list) else [dados]) if isinstance(d, dict)]}
    return {aba: list(csv.DictReader(texto.splitlines()))}


def valor_da_celula(valor: Any) -> Any:
    """Número que veio como texto vira número; senão a planilha não soma."""
    if not isinstance(valor, str):
        return valor
    s = valor.strip()
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        # "1.240" em tabela brasileira é milhar, não decimal.
        return int(s.replace(".", ""))
    if _NUMERO.match(s):
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            return float(s) if "." in s else int(s)
        except ValueError:
            return valor
    return valor


def montar(abas: Dict[str, List[Dict[str, Any]]]) -> Workbook:
    livro = Workbook()
    livro.remove(livro.active)
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    usados = set()
    for nome, linhas in abas.items():
        # Nome de aba no Excel: até 31 caracteres, sem []:*?/\ e sem repetir.
        base = re.sub(r"[\[\]:*?/\\]", " ", str(nome)).strip()[:31] or "Dados"
        titulo, n = base, 2
        while titulo.lower() in usados:
            titulo = f"{base[:28]} {n}"
            n += 1
        usados.add(titulo.lower())
        folha = livro.create_sheet(titulo)
        colunas: List[str] = []
        for linha in linhas:
            colunas += [c for c in linha if c not in colunas]
        folha.append(colunas)
        for linha in linhas:
            folha.append([valor_da_celula(linha.get(c)) for c in colunas])
        for celula in folha[1]:
            celula.font, celula.fill = cabecalho_fonte, cabecalho_fundo
            celula.alignment = Alignment(vertical="center")
        for i, coluna in enumerate(colunas, start=1):
            maior = max([len(str(coluna))] + [len(str(l.get(coluna) or "")) for l in linhas[:500]])
            folha.column_dimensions[get_column_letter(i)].width = min(max(maior + 2, 8), 60)
        if colunas:
            folha.freeze_panes = "A2"
            folha.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(linhas) + 1}"
    return livro


def main(argv: List[str]) -> int:
    parser = argparse.ArgumentParser(description="Converte CSV ou JSON numa planilha xlsx formatada.")
    parser.add_argument("entrada", help="arquivo .csv ou .json")
    parser.add_argument("saida", help="caminho do .xlsx a gravar")
    parser.add_argument("--aba", default="Dados", help="nome da aba quando a entrada tem uma tabela só")
    args = parser.parse_args(argv)

    abas = ler_abas(Path(args.entrada), args.aba)
    if not any(abas.values()):
        raise SystemExit("a entrada não tem linhas de dados")
    destino = Path(args.saida)
    destino.parent.mkdir(parents=True, exist_ok=True)
    montar(abas).save(destino)
    total = sum(len(v) for v in abas.values())
    print(f"Planilha gravada em {destino.resolve()} ({len(abas)} aba(s), {total} linhas)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

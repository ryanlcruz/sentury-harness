"""Roteiro em JSON para apresentação .pptx, sem escrever código de slide.

    python roteiro_para_pptx.py roteiro.json apresentacao.pptx

O roteiro tem `titulo`, `subtitulo` (opcional) e `slides`, uma lista em que
cada item é um destes:

    {"titulo": "...", "topicos": ["...", "..."], "notas": "..."}
    {"titulo": "...", "imagem": "grafico.png", "legenda": "..."}
    {"titulo": "...", "tabela": [["Coluna", "Valor"], ["a", "1"]]}
    {"titulo": "...", "colunas": [["à esquerda"], ["à direita"]]}
    {"secao": "Nome da parte"}

Sai em 16:9. Lista com mais de 7 tópicos vira mais de um slide, para o texto
não encolher até ninguém ler. Precisa do python-pptx.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

_MAX_TOPICOS = 7
_LARGURA, _ALTURA = Inches(13.333), Inches(7.5)
_MARGEM = Inches(0.6)
_TOPO_DO_CONTEUDO = Inches(1.6)

# Layouts do modelo padrão do python-pptx.
_CAPA, _TITULO_E_CONTEUDO, _SECAO, _SO_TITULO = 0, 1, 2, 5


def _titulo(slide, texto: str) -> None:
    if slide.shapes.title is not None:
        slide.shapes.title.text = texto


def _notas(slide, texto: str) -> None:
    if texto:
        slide.notes_slide.notes_text_frame.text = texto


def _area() -> tuple:
    return _MARGEM, _TOPO_DO_CONTEUDO, _LARGURA - 2 * _MARGEM, _ALTURA - _TOPO_DO_CONTEUDO - _MARGEM


def _topicos(apresentacao, item: Dict[str, Any]) -> None:
    topicos = [str(t) for t in item.get("topicos") or []]
    partes = [topicos[i:i + _MAX_TOPICOS] for i in range(0, len(topicos), _MAX_TOPICOS)] or [[]]
    for n, parte in enumerate(partes):
        slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[_TITULO_E_CONTEUDO])
        titulo = str(item.get("titulo") or "")
        _titulo(slide, titulo if n == 0 else f"{titulo} (continuação)")
        corpo = slide.placeholders[1].text_frame
        corpo.clear()
        for i, texto in enumerate(parte):
            paragrafo = corpo.paragraphs[0] if i == 0 else corpo.add_paragraph()
            paragrafo.text = texto
        _notas(slide, str(item.get("notas") or ""))


def _imagem(apresentacao, item: Dict[str, Any], base: Path) -> None:
    from PIL import Image

    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[_SO_TITULO])
    _titulo(slide, str(item.get("titulo") or ""))
    caminho = Path(str(item["imagem"]))
    if not caminho.is_absolute():
        caminho = base / caminho
    esquerda, topo, largura, altura = _area()
    if item.get("legenda"):
        altura -= Inches(0.5)
    # Cabe na área sem distorcer: a menor das duas escalas manda.
    with Image.open(caminho) as img:
        w, h = img.size
    escala = min(largura / w, altura / h)
    lw, lh = int(w * escala), int(h * escala)
    slide.shapes.add_picture(str(caminho), Emu(esquerda + (largura - lw) // 2), Emu(topo), Emu(lw), Emu(lh))
    if item.get("legenda"):
        caixa = slide.shapes.add_textbox(esquerda, Emu(topo + lh + Inches(0.1)), largura, Inches(0.4))
        caixa.text_frame.text = str(item["legenda"])
        caixa.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    _notas(slide, str(item.get("notas") or ""))


def _tabela(apresentacao, item: Dict[str, Any]) -> None:
    linhas: List[List[Any]] = [list(l) for l in item.get("tabela") or [] if isinstance(l, list)]
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[_SO_TITULO])
    _titulo(slide, str(item.get("titulo") or ""))
    if not linhas:
        return
    colunas = max(len(l) for l in linhas)
    esquerda, topo, largura, _ = _area()
    altura_linha = Inches(0.45)
    forma = slide.shapes.add_table(len(linhas), colunas, esquerda, topo, largura, altura_linha * len(linhas))
    for r, linha in enumerate(linhas):
        for c in range(colunas):
            celula = forma.table.cell(r, c)
            celula.text = "" if c >= len(linha) else str(linha[c])
            for paragrafo in celula.text_frame.paragraphs:
                for trecho in paragrafo.runs:
                    trecho.font.size = Pt(14)
                    trecho.font.bold = r == 0
    _notas(slide, str(item.get("notas") or ""))


def _colunas(apresentacao, item: Dict[str, Any]) -> None:
    lados = [list(c) if isinstance(c, list) else [str(c)] for c in (item.get("colunas") or [])][:2]
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[_SO_TITULO])
    _titulo(slide, str(item.get("titulo") or ""))
    esquerda, topo, largura, altura = _area()
    meia = (largura - Inches(0.4)) // 2
    for i, lado in enumerate(lados):
        caixa = slide.shapes.add_textbox(Emu(esquerda + i * (meia + Inches(0.4))), topo, Emu(meia), altura)
        quadro = caixa.text_frame
        quadro.word_wrap = True
        for j, texto in enumerate(lado):
            paragrafo = quadro.paragraphs[0] if j == 0 else quadro.add_paragraph()
            paragrafo.text = str(texto)
            paragrafo.runs[0].font.size = Pt(20)
    _notas(slide, str(item.get("notas") or ""))


def montar(roteiro: Dict[str, Any], base: Path) -> Presentation:
    apresentacao = Presentation()
    apresentacao.slide_width, apresentacao.slide_height = _LARGURA, _ALTURA
    capa = apresentacao.slides.add_slide(apresentacao.slide_layouts[_CAPA])
    _titulo(capa, str(roteiro.get("titulo") or "Apresentação"))
    if len(capa.placeholders) > 1:
        capa.placeholders[1].text = str(roteiro.get("subtitulo") or "")
    for item in roteiro.get("slides") or []:
        if not isinstance(item, dict):
            continue
        if "secao" in item:
            slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[_SECAO])
            _titulo(slide, str(item["secao"]))
        elif "imagem" in item:
            _imagem(apresentacao, item, base)
        elif "tabela" in item:
            _tabela(apresentacao, item)
        elif "colunas" in item:
            _colunas(apresentacao, item)
        else:
            _topicos(apresentacao, item)
    return apresentacao


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("roteiro", type=Path, help="roteiro em JSON")
    parser.add_argument("saida", type=Path, help="arquivo .pptx de saída")
    args = parser.parse_args(argv)
    try:
        roteiro = json.loads(args.roteiro.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"Não consegui ler o roteiro: {exc}", file=sys.stderr)
        return 1
    apresentacao = montar(roteiro, args.roteiro.resolve().parent)
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    apresentacao.save(str(args.saida))
    print(f"{args.saida} ({len(apresentacao.slides)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
